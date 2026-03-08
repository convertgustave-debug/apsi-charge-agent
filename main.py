from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pathlib import Path
from datetime import date, datetime
import pandas as pd
import numpy as np
from openpyxl.styles import PatternFill
from openpyxl.chart import BarChart, Reference
import requests

# =========================================================
# APP & DOSSIERS
# =========================================================

app = FastAPI()

UPLOAD_DIR = Path("uploads")
EXPORT_DIR = Path("exports")

UPLOAD_DIR.mkdir(exist_ok=True)
EXPORT_DIR.mkdir(exist_ok=True)

# =========================================================
# CONFIG CAPACITÉS
# =========================================================

class ChargeConfig(BaseModel):
    charge_max_1m: float = 60
    charge_max_3m: float = 130
    charge_max_6m: float = 210

CHARGE_CONFIG = ChargeConfig()

# =========================================================
# UTILS
# =========================================================

def normalize_str(x):
    if pd.isna(x):
        return None
    return str(x).strip().lower()

def parse_date_safe(x):
    try:
        d = pd.to_datetime(x, errors="coerce")
        return d.date() if not pd.isna(d) else None
    except:
        return None

def days_until(d):
    if d is None:
        return None
    return (d - date.today()).days

# =========================================================
# REGLES METIER
# =========================================================

SEGMENTATION_MAP = {
    "stratégique": 4,
    "strategique": 4,
    "projet": 3,
    "réassort": 2,
    "reassort": 2,
    "à développer": 1,
    "a développer": 1,
    "a developper": 1,
}

TRANSFO_MAP = {
    20: 0.70,
    40: 0.85,
    60: 1.00,
    80: 1.15,
    0.2: 0.70,
    0.4: 0.85,
    0.6: 1.00,
    0.8: 1.15,
    "20%": 0.70,
    "40%": 0.85,
    "60%": 1.00,
    "80%": 1.15,
}

AVG_COMPLEXITE = 2.5
AVG_SEGMENTATION = 2.5
AVG_TRANSFO = 0.925

def get_complexite(v):
    try:
        v = float(v)
        if 1 <= v <= 4:
            return v
        return AVG_COMPLEXITE
    except:
        return AVG_COMPLEXITE

def get_segmentation(v):
    if pd.isna(v):
        return AVG_SEGMENTATION
    return SEGMENTATION_MAP.get(normalize_str(v), AVG_SEGMENTATION)

def get_transfo(v):
    if pd.isna(v):
        return AVG_TRANSFO
    return TRANSFO_MAP.get(v, AVG_TRANSFO)

def coeff_urgence(days_left, horizon):

    if days_left is None:
        return 1.0

    if days_left < 0:
        return 1.5

    if horizon == 30:
        if days_left <= 7:
            return 1.4
        elif days_left <= 14:
            return 1.25
        else:
            return 1.1

    if horizon == 90:
        if days_left <= 15:
            return 1.5
        elif days_left <= 30:
            return 1.35
        elif days_left <= 60:
            return 1.15
        else:
            return 1.0

    if horizon == 180:
        if days_left <= 30:
            return 1.15
        elif days_left <= 60:
            return 1.30
        elif days_left <= 120:
            return 1.10
        else:
            return 1.0

    return 1.0

def coeff_ca(ca, ca_max):

    try:

        if pd.isna(ca):
            return 1.0

        if ca_max is None or ca_max <= 0:
            return 1.0

        return 1 + 0.1 * (float(ca) / ca_max)

    except:
        return 1.0

def compute_charge(row, ca_max, horizon):

    score = (
        0.5 * get_complexite(row["complexite"])
        + 0.5 * get_segmentation(row["segmentation"])
    )

    return (
        score
        * get_transfo(row["transformation"])
        * coeff_urgence(days_until(row["date_echeance"]), horizon)
        * coeff_ca(row["ca"], ca_max)
    )

# =========================================================
# EXPORT EXCEL
# =========================================================

def export_excel(df_detail, synthese_par_horizon):

    today = datetime.now().strftime("%Y-%m-%d")
    filename = f"Charge_CDP_{today}.xlsx"

    file_path = EXPORT_DIR / filename

    synthese = {}

    for horizon, rows in synthese_par_horizon.items():

        for r in rows:

            cdp = r["cdp"]

            if cdp not in synthese:
                synthese[cdp] = {"CDP": cdp}

            synthese[cdp][f"Charge {horizon}"] = r["charge_cdp"]
            synthese[cdp][f"Taux {horizon}"] = r["taux_charge_%"]

    df_synthese = pd.DataFrame(list(synthese.values()))

    numeric_cols_synth = df_synthese.select_dtypes(include=[np.number]).columns
    df_synthese[numeric_cols_synth] = df_synthese[numeric_cols_synth].round(1)

    numeric_cols_detail = df_detail.select_dtypes(include=[np.number]).columns
    df_detail[numeric_cols_detail] = df_detail[numeric_cols_detail].round(1)

    if "Charge 1M" in df_synthese.columns:
        df_synthese = df_synthese.sort_values("Charge 1M", ascending=False)

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:

        capacite_row = {
            "CDP": "Capacité max",
            "Charge 1M": 60,
            "Taux 1M": None,
            "Charge 3M": 130,
            "Taux 3M": None,
            "Charge 6M": 210,
            "Taux 6M": None
        }

        df_synthese = pd.concat(
            [df_synthese, pd.DataFrame([capacite_row])],
            ignore_index=True
        )

        df_synthese.to_excel(writer, sheet_name="Synthese_CDP", index=False)
        df_detail.to_excel(writer, sheet_name="Detail_Projets", index=False)

        worksheet = writer.sheets["Synthese_CDP"]

        # Colonne capacité restante (calculée dans Excel)
        worksheet.cell(row=1, column=8).value = "Capacité restante 1M"

        for row in range(2, worksheet.max_row):
            charge_cell = f"B{row}"
            worksheet.cell(row=row, column=8).value = f"=60-{charge_cell}"

        for row in range(2, worksheet.max_row):
            worksheet.cell(row=row, column=2).number_format = '0.0 "pts"'
            worksheet.cell(row=row, column=4).number_format = '0.0 "pts"'
            worksheet.cell(row=row, column=6).number_format = '0.0 "pts"'

            worksheet.cell(row=row, column=3).number_format = '0.0"%"'
            worksheet.cell(row=row, column=5).number_format = '0.0"%"'
            worksheet.cell(row=row, column=7).number_format = '0.0"%"'

        chart = BarChart()
        chart.title = "Charge CDP - Horizon 1M"
        chart.y_axis.title = "Points de charge"
        chart.x_axis.title = "CDP"

        data = Reference(worksheet, min_col=2, min_row=1, max_row=worksheet.max_row-1)
        cats = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row-1)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        worksheet.add_chart(chart, "J2")


        chart4 = BarChart()
        chart4.title = "Capacité restante par CDP (1M)"
        chart4.y_axis.title = "Points disponibles"
        chart4.x_axis.title = "CDP"

        data = Reference(
        worksheet,
         min_col=8,
            min_row=1,
            max_row=worksheet.max_row-1
        )

        cats = Reference(
            worksheet,
            min_col=1,
            min_row=2,
            max_row=worksheet.max_row-1
        )

        chart4.add_data(data, titles_from_data=True)
        chart4.set_categories(cats)

        worksheet.add_chart(chart4, "B20")

        chart2 = BarChart()
        chart2.title = "Comparaison des taux de charge"
        chart2.y_axis.title = "%"
        chart2.x_axis.title = "CDP"
    
        data1 = Reference(worksheet, min_col=3, min_row=1, max_row=worksheet.max_row-1)
        data2 = Reference(worksheet, min_col=5, min_row=1, max_row=worksheet.max_row-1)
        data3 = Reference(worksheet, min_col=7, min_row=1, max_row=worksheet.max_row-1)

        cats = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row-1)

        chart2.add_data(data1, titles_from_data=True)
        chart2.add_data(data2, titles_from_data=True)
        chart2.add_data(data3, titles_from_data=True)

        chart2.set_categories(cats)

        worksheet.add_chart(chart2, "J16")

        surcharge_rows = []

        for row in range(2, worksheet.max_row):

            taux = worksheet.cell(row=row, column=3).value

            try:
                if taux > 100:
                    surcharge_rows.append(row)
            except:
                pass

        if surcharge_rows:

            chart3 = BarChart()
            chart3.title = "CDP en surcharge (>100%)"
            chart3.y_axis.title = "%"
            chart3.x_axis.title = "CDP"

            data = Reference(
                worksheet,
                min_col=3,
                min_row=min(surcharge_rows),
                max_row=max(surcharge_rows)
            )

            cats = Reference(
                worksheet,
                min_col=1,
                min_row=min(surcharge_rows),
                max_row=max(surcharge_rows)
            )

            chart3.add_data(data, titles_from_data=False)
            chart3.set_categories(cats)

            worksheet.add_chart(chart3, "J29")

        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        orange_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        red_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")

        for row in range(2, worksheet.max_row):

            for col in [3,5,7]:

                cell = worksheet.cell(row=row, column=col)

                try:
                    value = float(cell.value)

                    if value < 70:
                        cell.fill = green_fill
                    elif value <= 100:
                        cell.fill = orange_fill
                    else:
                        cell.fill = red_fill

                except:
                    pass

        grey_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

        for col in range(1, 8):
            worksheet.cell(row=worksheet.max_row, column=col).fill = grey_fill

        worksheet.freeze_panes = "A2"

        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 25)

    return file_path

# =========================================================
# ENDPOINT API
# =========================================================

@app.post("/process")
async def process_file(payload: dict):

    try:

        file_url = payload.get("file_url")

        if not file_url:
            raise HTTPException(400, "file_url manquant")

        r = requests.get(file_url)

        if r.status_code != 200:
            raise HTTPException(400, "Impossible de télécharger le fichier")

        input_path = UPLOAD_DIR / "input.xlsx"

        with open(input_path, "wb") as f:
            f.write(r.content)

        df = pd.read_excel(input_path, engine="openpyxl")

        df.columns = [c.strip().lower() for c in df.columns]

        col_map = {
            "cdp": ["cdp", "cdp mobilier"],
            "statut": ["statut de l'opportunité", "statut"],
            "date_echeance": ["date d'échéance du projet", "échéance", "echéance opport mob"],
            "complexite": ["complexité", "complexité du projet"],
            "segmentation": ["segmentation", "segmentation mob"],
            "transformation": ["tx de transfo", "tx de transfo mob"],
            "ca": ["ca potentiel mob"],
        }

        def pick(cols):
            for c in cols:
                if c in df.columns:
                    return c
            return None

        cols = {k: pick(v) for k, v in col_map.items()}

        if any(v is None for v in cols.values()):
            raise HTTPException(400, f"Colonnes manquantes : {cols}")

        work = pd.DataFrame({
            "cdp": df[cols["cdp"]],
            "statut": df[cols["statut"]],
            "date_echeance": df[cols["date_echeance"]].apply(parse_date_safe),
            "complexite": df[cols["complexite"]],
            "segmentation": df[cols["segmentation"]],
            "transformation": df[cols["transformation"]],
            "ca": df[cols["ca"]],
        })

        work["statut_norm"] = work["statut"].apply(normalize_str)

        devis = work[work["statut_norm"] == "devis en cours"].copy()

        ca_max = pd.to_numeric(devis["ca"], errors="coerce").max()
        ca_max = None if pd.isna(ca_max) else float(ca_max)

        horizons = [
            ("1M", 30, CHARGE_CONFIG.charge_max_1m),
            ("3M", 90, CHARGE_CONFIG.charge_max_3m),
            ("6M", 180, CHARGE_CONFIG.charge_max_6m),
        ]

        result = {}

        for label, days, cap in horizons:

            tmp = devis.copy()

            tmp["charge_projet"] = tmp.apply(
                lambda r: compute_charge(r, ca_max, days),
                axis=1
            )

            agg = tmp.groupby("cdp")["charge_projet"].sum().reset_index()

            agg["taux_charge_%"] = agg["charge_projet"] / cap * 100

            result[label] = [
                {
                    "cdp": r["cdp"],
                    "charge_cdp": round(float(r["charge_projet"]),1),
                    "taux_charge_%": round(float(r["taux_charge_%"]),1),
                }
                for _, r in agg.sort_values("charge_projet", ascending=False).iterrows()
            ]

        output_path = export_excel(devis, result)

        return FileResponse(
            path=output_path,
            filename=output_path.name,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        raise HTTPException(500, str(e))


















